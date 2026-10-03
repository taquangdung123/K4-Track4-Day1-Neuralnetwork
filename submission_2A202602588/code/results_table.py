"""Serialization and spreadsheet utilities for experiment results.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import torch


EXPERIMENT_COLUMNS = (
    "exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay",
    "batch", "epochs", "hidden", "dropout", "clip_norm", "precision", "init",
    "seed", "step0_loss", "best_val_loss", "best_epoch", "final_train_loss",
    "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB",
    "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes",
)
FORMULA_COLUMNS = {
    "step0_gap_vs_lnC",
    "gap_val_minus_train",
    "delta_val_f1_vs_base",
    "beyond_noise",
}


def _json_value(value):
    if isinstance(value, dict):
        return {str(key): _json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, torch.Tensor):
        return value.detach().cpu().tolist()
    if isinstance(value, Path):
        return str(value)
    return value


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    required = ("cfg", "history", "summary")
    missing = [key for key in required if key not in result]
    if missing:
        raise ValueError(f"Result is missing required sections: {', '.join(missing)}")
    exp_id = result["cfg"].get("exp_id")
    if not exp_id or Path(str(exp_id)).name != str(exp_id):
        raise ValueError("cfg['exp_id'] must be a non-empty filename-safe identifier")
    output_dir = Path(results_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = {key: _json_value(result[key]) for key in required}
    output_path = output_dir / f"{exp_id}.json"
    with output_path.open("w", encoding="utf-8") as output_file:
        json.dump(payload, output_file, indent=2, allow_nan=False)
    return str(output_path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    directory = Path(results_dir)
    if not directory.is_dir():
        raise FileNotFoundError(f"Results directory not found: {directory.resolve()}")
    loaded = []
    for result_path in sorted(directory.glob("*.json")):
        with result_path.open(encoding="utf-8") as result_file:
            result = json.load(result_file)
        result.setdefault("cfg", {}).setdefault("exp_id", result_path.stem)
        loaded.append(result)
    return sorted(loaded, key=lambda result: result["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg = result.get("cfg", {})
    summary = result.get("summary", {})
    exp_id = cfg.get("exp_id")
    if not exp_id:
        raise ValueError("result.cfg must contain exp_id")
    row = {column: "" for column in EXPERIMENT_COLUMNS}
    for key in (
        "exp_id", "group", "description", "loss", "optimizer", "lr",
        "weight_decay", "batch", "epochs", "dropout", "clip_norm",
        "precision", "init", "seed",
    ):
        row[key] = cfg.get(key, "")
    row["hidden"] = str(tuple(cfg["hidden"])) if cfg.get("hidden") is not None else ""
    for key in (
        "step0_loss", "best_val_loss", "best_epoch", "final_train_loss",
        "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s",
        "peak_mem_MB", "diverged",
    ):
        row[key] = summary.get(key, "")
    if eval_scores is not None:
        row["eval_acc"] = eval_scores.get("accuracy", "")
        row["eval_macro_f1"] = eval_scores.get("macro_f1", "")
    row["figure_file"] = f"figures/{exp_id}.png"
    row["notes"] = notes
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    try:
        from copy import copy
        from openpyxl import load_workbook
        from openpyxl.formula.translate import Translator
    except ImportError as error:
        raise ImportError("write_xlsx requires openpyxl; install the project's spreadsheet dependencies") from error

    template = Path(template_path)
    if not template.is_file():
        raise FileNotFoundError(f"Spreadsheet template not found: {template.resolve()}")
    workbook = load_workbook(template)
    if "Experiments" not in workbook.sheetnames:
        raise KeyError("Template workbook does not contain an 'Experiments' sheet")
    worksheet = workbook["Experiments"]
    headers = {
        str(cell.value): cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }
    missing_headers = set(EXPERIMENT_COLUMNS) - headers.keys()
    if missing_headers:
        raise ValueError(
            "Template is missing required columns: " + ", ".join(sorted(missing_headers))
        )

    formula_columns = FORMULA_COLUMNS.intersection(headers)
    first_data_row = 2
    template_row = first_data_row
    for row_index, row in enumerate(rows, start=first_data_row):
        if row_index > worksheet.max_row:
            worksheet.row_dimensions[row_index].height = worksheet.row_dimensions[template_row].height
            for column_index in range(1, worksheet.max_column + 1):
                source = worksheet.cell(template_row, column_index)
                destination = worksheet.cell(row_index, column_index)
                if source.has_style:
                    destination._style = copy(source._style)
                if source.number_format:
                    destination.number_format = source.number_format
                if source.value is not None and source.data_type == "f":
                    destination.value = Translator(
                        source.value, origin=source.coordinate
                    ).translate_formula(destination.coordinate)

        for header, column_index in headers.items():
            if header in formula_columns:
                continue
            worksheet.cell(row_index, column_index).value = row.get(header, "")

    output = Path(out_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output)
